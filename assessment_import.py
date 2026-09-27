from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook


def txt(v: Any) -> str:
    return "" if v is None else str(v).strip()


def regno(v: Any) -> str:
    s=txt(v)
    return s[:-2] if s.endswith(".0") else s


def num(v: Any) -> float:
    try: return float(v)
    except (TypeError,ValueError): return 0.0


def qkey(v: Any) -> str:
    key=re.sub(r"[^A-Z0-9]","",txt(v).upper()).removeprefix("QUESTION")
    return key[1:] if key.startswith("Q") and len(key)>1 and key[1].isdigit() else key


@dataclass
class Question:
    key: str
    maximum: float
    co: str
    bloom: str=""


@dataclass
class Assessment:
    name: str
    kind: str
    maximum: float
    questions: list[Question]


def parse_qp_analysis(path: Path) -> list[Assessment]:
    if path.suffix.lower() != ".docx":
        raise ValueError("QP Analysis auto-reading currently requires a .docx file.")
    doc=Document(path)
    summary=[]
    for table in doc.tables:
        rows=[[txt(c.text) for c in row.cells] for row in table.rows]
        if rows and any("name of the assessment" in x.lower() for x in rows[0]):
            for row in rows[1:]:
                if len(row)>=3 and row[1] and "total" not in row[1].lower(): summary.append((row[1],num(row[2])))
            break
    maps=[]
    for table in doc.tables:
        rows=[[txt(c.text) for c in row.cells] for row in table.rows]
        if not rows: continue
        header_row=next((i for i,row in enumerate(rows[:3]) if any("mark" in x.lower() for x in row) and any(re.fullmatch(r"\s*co(?:\*|/clo|\s*/clo)?\s*",x,re.I) for x in row)),0)
        head=[x.lower() for x in rows[header_row]]
        qi=next((i for i,x in enumerate(head) if "q." in x or "q.no" in x or "q no" in x),None)
        mi=next((i for i,x in enumerate(head) if "mark" in x),None)
        ci=next((i for i,x in enumerate(head) if x in {"co","co*","co/clo","co /clo"} or "co/clo" in x),None)
        bi=next((i for i,x in enumerate(head) if "bloom" in x or x=="bl"),None)
        if qi is None or mi is None or ci is None: continue
        questions=[]
        for row in rows[header_row+1:]:
            if max(qi,mi,ci)>=len(row): continue
            co_match=re.search(r"\bCO\s*(\d+)\b",row[ci],re.I)
            if not co_match or not qkey(row[qi]): continue
            questions.append(Question(qkey(row[qi]),num(row[mi]),f"CO{co_match.group(1)}",row[bi] if bi is not None and bi<len(row) else ""))
        if questions: maps.append(questions)
    if not maps: raise ValueError("No question-to-CO mapping tables were found in the QP Analysis document.")
    assessments=[]
    for i,questions in enumerate(maps):
        name,maximum=summary[i] if i<len(summary) else (f"Assessment {i+1}",sum(q.maximum for q in questions))
        low=name.lower(); kind="mid" if "mid" in low else "final" if "end" in low or "final" in low or "see" in low else "ia"
        assessments.append(Assessment(name,kind,maximum or sum(q.maximum for q in questions),questions))
    return assessments


def read_mark_table(path: Path) -> tuple[list[str],list[dict[str,Any]]]:
    if path.suffix.lower()==".csv":
        with path.open(encoding="utf-8-sig",newline="") as f: rows=list(csv.reader(f))
    else:
        wb=load_workbook(path,data_only=True,read_only=False)
        ws=next((s for s in wb.worksheets if s.max_row and s.max_column),wb.active)
        rows=[list(r) for r in ws.iter_rows(values_only=True)]
    header_i=-1
    for i,row in enumerate(rows):
        lowered=[txt(x).lower() for x in row]
        if any("registration" in x or "enrollment" in x or x in {"reg no","reg. no."} for x in lowered): header_i=i;break
    if header_i<0:
        # Common compact IA sheet: registration number, name, then assessment columns.
        for i,row in enumerate(rows):
            if row and re.fullmatch(r"\d{6,}",regno(row[0])): header_i=max(0,i-1);break
    if header_i<0: raise ValueError(f"Could not locate a registration-number column in {path.name}.")
    raw_header=rows[header_i]
    headers=[txt(x) or f"COL{i+1}" for i,x in enumerate(raw_header)]
    reg_i=next((i for i,x in enumerate(headers) if "registration" in x.lower() or "enrollment" in x.lower() or x.lower() in {"reg no","reg. no."}),0)
    name_i=next((i for i,x in enumerate(headers) if "name" in x.lower()),1)
    records=[]
    for row in rows[header_i+1:]:
        row=list(row)+[None]*(len(headers)-len(row)); reg=regno(row[reg_i])
        if not re.fullmatch(r"\d{6,}",reg): continue
        records.append({"reg":reg,"name":txt(row[name_i]),"values":{headers[i]:num(row[i]) for i in range(len(headers)) if i not in {reg_i,name_i}}})
    if not records: raise ValueError(f"No student records were found in {path.name}.")
    return headers,records


def match_question(header: str, questions: list[Question]) -> list[Question]:
    key=qkey(header)
    direct=[q for q in questions if q.key==key]
    if direct: return direct
    return [q for q in questions if re.match(rf"^{re.escape(key)}[A-Z]$",q.key)]


def preview_conflicts(mark_path: Path, assessment: Assessment) -> list[dict[str,Any]]:
    _,records=read_mark_table(mark_path); conflicts=[]
    headers=list(records[0]["values"])
    for header in headers:
        matched=match_question(header,assessment.questions)
        cos=sorted({q.co for q in matched})
        if len(cos)>1:
            conflicts.append({"assessment":assessment.name,"column":header,"cos":cos,"subparts":[{"question":q.key,"co":q.co,"maximum":q.maximum} for q in matched],"students":[{"reg":r["reg"],"name":r["name"],"mark":r["values"][header]} for r in records]})
    return conflicts


def _score_columns(records: list[dict[str,Any]]) -> list[str]:
    columns=list(records[0]["values"])
    rejected=("total","difference","percent","grade","sl.","serial")
    columns=[c for c in columns if not any(x in c.lower() for x in rejected)]
    if len(columns)>1:
        last=columns[-1]; prior=columns[:-1]
        if sum(abs(r["values"][last]-sum(r["values"][c] for c in prior))<0.001 for r in records)>=max(1,len(records)-1): columns=prior
    return columns


def normalize_assessments(qp_path: Path, ia_paths: list[Path], mid_path: Path, final_path: Path, co_labels: list[str], manual_splits: list[dict[str,Any]]) -> tuple[list[dict[str,Any]],list[float],list[float],list[float],list[str]]:
    assessments=parse_qp_analysis(qp_path); ia_maps=[a for a in assessments if a.kind=="ia"]; mid_map=next((a for a in assessments if a.kind=="mid"),None); final_map=next((a for a in assessments if a.kind=="final"),None)
    if not mid_map or not final_map: raise ValueError("QP Analysis must contain Midsem and final-examination question mappings.")
    index={co:i for i,co in enumerate(co_labels)}; students={}; warnings=[]
    def student(reg,name):
        row=students.setdefault(reg,{"reg":reg,"name":name,"ia":[0.0]*len(co_labels),"mid":[0.0]*len(co_labels),"see":[0.0]*len(co_labels)})
        if name and not row["name"]: row["name"]=name
        return row
    ia_max=[0.0]*len(co_labels);mid_max=[0.0]*len(co_labels);see_max=[0.0]*len(co_labels)
    # Combined IA file: one score column per IA; separate files: one file per IA.
    ia_cursor=0
    for path in ia_paths:
        _,records=read_mark_table(path); cols=_score_columns(records)
        maps=ia_maps[ia_cursor:ia_cursor+len(cols)] if len(ia_paths)==1 else ia_maps[ia_cursor:ia_cursor+1]
        if len(ia_paths)>1: cols=cols[:1]
        if len(maps)!=len(cols): raise ValueError(f"Could not match IA columns in {path.name} to the QP Analysis assessments.")
        for col,assessment in zip(cols,maps):
            cos={q.co for q in assessment.questions};
            if len(cos)!=1: raise ValueError(f"{assessment.name} contains multiple COs; upload question-wise IA marks.")
            co=next(iter(cos)); ci=index.get(co)
            if ci is None: continue
            ia_max[ci]+=assessment.maximum
            for record in records: student(record["reg"],record["name"])["ia"][ci]+=record["values"][col]
        ia_cursor+=len(maps)
    if ia_cursor<len(ia_maps): warnings.append(f"Only {ia_cursor} of {len(ia_maps)} internal assessments were matched.")
    split_index={(x["assessment"],qkey(x["column"]),x["reg"]):x for x in manual_splits}
    def apply(path,assessment,bucket,maxima):
        _,records=read_mark_table(path);cols=_score_columns(records)
        for col in cols:
            questions=match_question(col,assessment.questions)
            if not questions: continue
            by_co={}
            for q in questions: by_co[q.co]=by_co.get(q.co,0)+q.maximum
            for co,mx in by_co.items():
                if co in index: maxima[index[co]]+=mx
            for record in records:
                target=student(record["reg"],record["name"])[bucket]
                if len(by_co)==1: target[index[next(iter(by_co))]]+=record["values"][col]
                else:
                    split=split_index.get((assessment.name,qkey(col),record["reg"]))
                    if not split: raise ValueError(f"Manual CO split is required for {assessment.name} {col}, registration {record['reg']}.")
                    total=sum(num(p.get("mark")) for p in split.get("parts",[]))
                    if abs(total-record["values"][col])>0.001: raise ValueError(f"Manual split for {record['reg']} {col} does not equal the original mark.")
                    for part in split["parts"]:
                        if part.get("co") in index: target[index[part["co"]]]+=num(part.get("mark"))
    apply(mid_path,mid_map,"mid",mid_max);apply(final_path,final_map,"see",see_max)
    ordered=sorted(students.values(),key=lambda x:(0,int(x["reg"])) if x["reg"].isdigit() else (1,x["reg"]))
    return ordered,ia_max,mid_max,see_max,warnings
