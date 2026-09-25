from __future__ import annotations

import json
import math
import re
import shutil
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font

from xlsx_raw import first_nonempty_sheet, read_xlsx


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def number(value: Any, default: float = 0.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def canonical_reg(value: Any) -> str:
    raw = text(value)
    return raw[:-2] if raw.endswith(".0") else raw


@dataclass
class Student:
    reg: str
    name: str
    ia: list[float]
    mid: list[float]
    see: list[float]
    grade: str = ""


@dataclass
class Course:
    code: str = ""
    name: str = ""
    faculty: str = ""
    school: str = "Department of Biomedical Engineering"
    program: str = "M.Tech Biomedical Engineering (Medical Informatics)"
    semester: str = "II"
    year: str = ""
    odd_even: str = "Even"
    section: str = "Medical Informatics"
    course_type: str = "Core"
    cos: list[str] = field(default_factory=lambda: ["", "", "", ""])
    bloom: list[str] = field(default_factory=lambda: ["L3", "L3", "L5", "L6"])


DEFAULT_MAPPING = [[2, 1, 3, 0, 0, 0], [3, 1, 3, 0, 0, 0], [3, 2, 3, 0, 0, 0], [3, 3, 3, 0, 0, 0]]


def parse_marks(path: Path) -> tuple[list[Student], list[float], list[float], list[float]]:
    sheets = read_xlsx(path)
    rows = sheets.get("CO Summary") or next(iter(sheets.values()), [])
    header_index = next((i for i, row in enumerate(rows) if len(row) > 1 and text(row[0]).lower() == "registration no."), -1)
    if header_index < 0 or header_index + 1 >= len(rows):
        raise ValueError("Marks file must contain a 'CO Summary' table with Registration No. and CO1–CO4 columns.")
    maximum = rows[header_index + 1]
    ia_max = [number(x) for x in maximum[2:6]]
    mid_max = [number(x) for x in maximum[7:11]]
    see_max = [number(x) for x in maximum[12:16]]
    students: list[Student] = []
    for row in rows[header_index + 2:]:
        reg = canonical_reg(row[0] if row else None)
        if not reg or not re.fullmatch(r"\d{6,}", reg):
            continue
        padded = list(row) + [None] * (17 - len(row))
        students.append(Student(reg, text(padded[1]), [number(x) for x in padded[2:6]], [number(x) for x in padded[7:11]], [number(x) for x in padded[12:16]]))
    if not students:
        raise ValueError("No student rows were found in the marks file.")
    return students, ia_max, mid_max, see_max


def apply_grades(path: Path | None, students: list[Student], course_code: str) -> list[str]:
    warnings: list[str] = []
    if not path:
        warnings.append("Grade file was not supplied; grade cells will remain blank.")
        return warnings
    rows = first_nonempty_sheet(path)
    if not rows:
        return ["Grade file was empty or unreadable."]
    header = [text(x).lower() for x in rows[0]]
    def locate(fragment: str, fallback: int) -> int:
        return next((i for i, item in enumerate(header) if fragment in item), fallback)
    code_i, name_i = locate("course code", 2), locate("participant account: name", 3)
    reg_i, grade_i = locate("enrollment id", 4), locate("grade", 10)
    by_reg: dict[str, tuple[str, str]] = {}
    for row in rows[1:]:
        padded = list(row) + [None] * (max(code_i, name_i, reg_i, grade_i) + 1 - len(row))
        if course_code and text(padded[code_i]).replace(" ", "").lower() != course_code.replace(" ", "").lower():
            continue
        reg = canonical_reg(padded[reg_i])
        if reg:
            by_reg[reg] = (text(padded[name_i]), text(padded[grade_i]))
    for student in students:
        if student.reg in by_reg:
            student.name = by_reg[student.reg][0].title()
            student.grade = by_reg[student.reg][1]
        else:
            warnings.append(f"No grade matched registration number {student.reg}.")
    return warnings


def parse_course_plan(path: Path | None, fallback: Course) -> Course:
    if not path:
        return fallback
    doc = Document(path)
    content = "\n".join([p.text for p in doc.paragraphs] + [" | ".join(text(c.text) for c in row.cells) for table in doc.tables for row in table.rows])
    patterns = {
        "code": r"Course\s*Code\s*[:|]\s*([A-Z]{2,}\s*\d{3,})",
        "name": r"Course\s*(?:Name|Title)\s*[:|]\s*([^\n|]+)",
    }
    for key, pattern in patterns.items():
        match = re.search(pattern, content, flags=re.I)
        if match and not getattr(fallback, key):
            setattr(fallback, key, match.group(1).strip())
    co_matches = re.findall(r"\bCO\s*([1-6])\s*[:.\-|]\s*([^\n|]+)", content, flags=re.I)
    for index, statement in co_matches:
        position = int(index) - 1
        if position < 4 and len(statement.strip()) > 15:
            fallback.cos[position] = statement.strip()
    return fallback


def parse_survey(path: Path | None) -> list[dict[str, Any]]:
    if not path:
        return []
    rows = first_nonempty_sheet(path)
    if not rows:
        return []
    header = [text(x).lower().replace("\xa0", " ") for x in rows[0]]
    name_i = next((i for i, x in enumerate(header) if "student's full name" in x or "student full name" in x), 6)
    reg_i = next((i for i, x in enumerate(header) if "reg. no" in x or "registration" in x), 7)
    co_indices = []
    for co in range(1, 5):
        pattern = re.compile(rf"\bco\s*{co}\b")
        co_indices.append(next((i for i, x in enumerate(header) if pattern.search(x)), 7 + co))
    responses = []
    for row in rows[1:]:
        padded = list(row) + [None] * (max([name_i, reg_i] + co_indices) + 1 - len(row))
        reg = canonical_reg(padded[reg_i])
        if not reg:
            continue
        ratings = [number(padded[i]) for i in co_indices]
        if any(ratings):
            responses.append({"name": text(padded[name_i]), "reg": reg, "ratings": ratings})
    return responses


def attainment(mark: float, maximum: float, grade: str) -> int:
    if maximum <= 0 or grade.upper() in {"DT", "I"}:
        return 0
    percentage = mark / maximum * 100
    return 3 if percentage >= 75 else 2 if percentage >= 50 else 1


def weighted_level(levels: list[int]) -> float:
    valid = [x for x in levels if x > 0]
    return sum(valid) / len(valid) if valid else 0.0


def compute(students: list[Student], ia_max: list[float], mid_max: list[float], see_max: list[float], responses: list[dict[str, Any]], direct_weights: tuple[float, float], overall_weights: tuple[float, float], mapping: list[list[float]], target: float) -> dict[str, Any]:
    cie_max = [ia_max[i] + mid_max[i] for i in range(4)]
    cie_marks = [[s.ia[i] + s.mid[i] for i in range(4)] for s in students]
    cie_levels = [[attainment(cie_marks[r][c], cie_max[c], students[r].grade) for c in range(4)] for r in range(len(students))]
    see_levels = [[attainment(students[r].see[c], see_max[c], students[r].grade) for c in range(4)] for r in range(len(students))]
    cie = [weighted_level([row[c] for row in cie_levels]) for c in range(4)]
    see = [weighted_level([row[c] for row in see_levels]) for c in range(4)]
    direct = [(cie[i] * direct_weights[0] + see[i] * direct_weights[1]) / 100 for i in range(4)]
    indirect = []
    for c in range(4):
        levels = [3 if x["ratings"][c] == 5 else 2 if x["ratings"][c] >= 3 else 1 for x in responses]
        indirect.append(weighted_level(levels))
    overall = [((direct[i] * overall_weights[0] + indirect[i] * overall_weights[1]) / 100) if responses else direct[i] for i in range(4)]
    def po(co_values: list[float]) -> list[float]:
        result = []
        for p in range(6):
            denominator = sum(mapping[c][p] for c in range(4))
            result.append(sum(co_values[c] * mapping[c][p] for c in range(4)) / denominator if denominator else 0.0)
        return result
    return {"cie_max": cie_max, "cie_marks": cie_marks, "cie_levels": cie_levels, "see_levels": see_levels, "cie": cie, "see": see, "direct": direct, "indirect": indirect, "overall": overall, "direct_po": po(direct), "overall_po": po(overall), "target": target}


def _sheet(workbook, title: str):
    normalized = title.strip()
    return next((ws for ws in workbook.worksheets if ws.title.strip() == normalized), None)


def fill_template(template: Path, output: Path, course: Course, students: list[Student], maxima: tuple[list[float], list[float], list[float]], survey: list[dict[str, Any]], result: dict[str, Any], mapping: list[list[float]], target: float, direct_weights: tuple[float, float], overall_weights: tuple[float, float]) -> None:
    shutil.copy2(template, output)
    wb = load_workbook(output)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    ia_max, mid_max, see_max = maxima

    ws = _sheet(wb, "2. TARGET")
    for cell, value in {"E3":course.school,"N3":course.program,"P3":course.year,"B4":course.semester,"E4":course.code,"I4":course.name,"N4":course.faculty,"N5":course.course_type}.items(): ws[cell] = value
    for i in range(4): ws.cell(18, 2+i, target); ws.cell(23, 2+i, target)

    ws = _sheet(wb, "3. CO-PO Mapping")
    ws["S4"], ws["S5"] = course.odd_even, course.section
    for i in range(4):
        ws.cell(13+i, 4, course.cos[i]); ws.cell(13+i, 21, course.bloom[i]); ws.cell(24+i, 1, f"CO{i+1}")
        for p in range(6): ws.cell(24+i, 4+p, mapping[i][p])

    ws = _sheet(wb, "4. CIE Assessment Marks")
    for row in range(14, 49):
        for col in range(2, 54): ws.cell(row, col, None)
    max_cells = ["D12","L12","T12","AB12"]
    for i, cell in enumerate(max_cells): ws[cell] = ia_max[i]
    for i, col in enumerate(range(32,36)): ws.cell(12,col,mid_max[i])
    total_cols = [46,47,48,49,50,51]
    for i in range(4): ws.cell(12,total_cols[i],ia_max[i]+mid_max[i])
    ws["AZ12"] = sum(ia_max)+sum(mid_max)
    for index, student in enumerate(students, start=14):
        ws.cell(index,2,student.name); ws.cell(index,3,student.reg); ws.cell(index,4,student.ia[0]); ws.cell(index,12,student.ia[1]); ws.cell(index,20,student.ia[2]); ws.cell(index,28,student.ia[3])
        for i in range(4): ws.cell(index,32+i,student.mid[i]); ws.cell(index,46+i,student.ia[i]+student.mid[i])
        ws.cell(index,52,sum(student.ia)+sum(student.mid)); ws.cell(index,53,student.grade)

    ws = _sheet(wb, "5. CIE- CO Attainment")
    for i in range(4): ws.cell(21,4+i,result["cie_max"][i])
    for r, student in enumerate(students, start=22):
        idx=r-22; ws.cell(r,2,student.name); ws.cell(r,3,student.reg); ws.cell(r,22,student.grade)
        for c in range(4):
            ws.cell(r,4+c,result["cie_marks"][idx][c]); ws.cell(r,10+2*c,result["cie_marks"][idx][c]/result["cie_max"][c]*100 if result["cie_max"][c] else 0); ws.cell(r,11+2*c,result["cie_levels"][idx][c])
    for c in range(4): ws.cell(15,9+c,result["cie"][c])
    ws["C57"] = len(students)

    ws = _sheet(wb, "6. SEE - Marks")
    starts=[4,9,14,19]
    for i,col in enumerate(starts): ws.cell(13,col,see_max[i])
    for r,student in enumerate(students,start=15):
        ws.cell(r,2,student.name); ws.cell(r,3,student.reg); ws.cell(r,41,student.grade)
        for c,col in enumerate(starts): ws.cell(r,col,student.see[c]); ws.cell(r,34+c,student.see[c])
        ws.cell(r,40,sum(student.see))

    ws = _sheet(wb, "7. SEE- CO Attainment")
    for i in range(4): ws.cell(21,4+i,see_max[i])
    for r,student in enumerate(students,start=22):
        idx=r-22; ws.cell(r,2,student.name); ws.cell(r,3,student.reg); ws.cell(r,22,student.grade)
        for c in range(4):
            ws.cell(r,4+c,student.see[c]); ws.cell(r,10+2*c,student.see[c]/see_max[c]*100 if see_max[c] else 0); ws.cell(r,11+2*c,result["see_levels"][idx][c])
    for c in range(4): ws.cell(15,9+c,result["see"][c])
    ws["C57"] = len(students)

    ws = _sheet(wb, "8. Course Feedback")
    for c in range(4): ws.cell(23,5+c,course.cos[c])
    for row in range(25,60):
        for col in range(2,15): ws.cell(row,col,None)
    for r,response in enumerate(survey,start=25):
        ws.cell(r,2,response["name"]); ws.cell(r,3,response["reg"])
        for c,rating in enumerate(response["ratings"]):
            ws.cell(r,5+c,rating); ws.cell(r,11+c,3 if rating==5 else 2 if rating>=3 else 1)
    for c in range(4):
        levels=[3 if x["ratings"][c]==5 else 2 if x["ratings"][c]>=3 else 1 for x in survey]
        ws.cell(12,11+c,levels.count(3)); ws.cell(13,11+c,levels.count(2)); ws.cell(14,11+c,levels.count(1)); ws.cell(16,11+c,len(survey)); ws.cell(17,11+c,len(students)); ws.cell(18,11+c,sum(levels)); ws.cell(19,11+c,result["indirect"][c] if survey else "NA")
    ws["A20"] = f"Course outcome survey response rate: {len(survey)} of {len(students)} students ({len(survey)/len(students)*100:.0f}%)." if students else "No enrolled students found."
    ws["A20"].fill=PatternFill("solid",fgColor="D9EAF7"); ws["A20"].font=Font(color="1F4E78",italic=True)

    ws = _sheet(wb, "9.Direct & Overall CO attinment")
    ws["B14"],ws["B15"] = direct_weights; ws["C19"],ws["C20"] = overall_weights
    for c in range(4):
        col=4+c; ws.cell(14,col,result["cie"][c]); ws.cell(15,col,result["see"][c]); ws.cell(16,col,result["direct"][c]); ws.cell(19,col,result["direct"][c]); ws.cell(20,col,result["indirect"][c] if survey else "NA"); ws.cell(21,col,result["overall"][c]); ws.cell(22,col,target); ws.cell(23,col,"YES" if result["overall"][c]>=target else "NO")

    ws = _sheet(wb, "10. ACTION PLAN & RCA-CO")
    actions=["Use healthcare-agent examples and short diagnostic quizzes.","Use guided cases on algorithm selection and heuristic problem solving.","Reinforce MLOps, data-integration and deployment-validation case studies.","Use ethics, safety, human-in-the-loop discussions and a capstone presentation rubric."]
    for i in range(4): ws.cell(8+i,8,actions[i])
    ws["A15"],ws["A18"] = course.faculty,"Review COs below target and retain successful assessment methods for attained COs."

    ws = _sheet(wb, "11. PO ATTAINMENT")
    for c in range(4): ws.cell(9+c,1,result["direct"][c]); ws.cell(22+c,1,result["overall"][c])
    for p in range(6):
        ws.cell(15,5+p,result["direct_po"][p]); ws.cell(16,5+p,target if result["direct_po"][p]>0 else "NA"); ws.cell(17,5+p,"YES" if result["direct_po"][p]>=target and result["direct_po"][p]>0 else "NA")
        ws.cell(28,5+p,result["overall_po"][p]); ws.cell(29,5+p,target if result["overall_po"][p]>0 else "NA"); ws.cell(30,5+p,"YES" if result["overall_po"][p]>=target and result["overall_po"][p]>0 else "NA")

    ws = _sheet(wb, "12. ACTION PLAN & RCA-PO")
    po_actions=["Continue clinical case investigations, evidence-based justification and capstone validation.","Strengthen technical reporting and presentation using structured rubrics.","Deepen specialization through advanced AI, MLOps and ethics integration."]
    for i in range(3): ws.cell(8+i,8,po_actions[i])
    ws["A15"],ws["A18"] = course.faculty,"Review mapped POs below target and document the corrective action in the next delivery cycle."
    wb.save(output)


def run_job(files: dict[str, Path | None], fields: dict[str, str], output_dir: Path) -> tuple[Path, dict[str, Any]]:
    course = Course(code=fields.get("course_code", ""), name=fields.get("course_name", ""), faculty=fields.get("faculty", ""), school=fields.get("school", "") or Course.school, program=fields.get("program", "") or Course.program, semester=fields.get("semester", "II"), year=fields.get("year", ""), odd_even=fields.get("odd_even", "Even"), section=fields.get("section", "Medical Informatics"), course_type=fields.get("course_type", "Core"))
    course = parse_course_plan(files.get("course_plan"), course)
    for i in range(4):
        if fields.get(f"co{i+1}"): course.cos[i]=fields[f"co{i+1}"]
    students, ia_max, mid_max, see_max = parse_marks(files["marks"])
    warnings = apply_grades(files.get("grades"), students, course.code)
    registrations=[s.reg for s in students]
    if len(registrations)!=len(set(registrations)): raise ValueError("Duplicate registration numbers were found in the marks file.")
    survey = parse_survey(files.get("survey"))
    if survey and len(survey) < len(students): warnings.append(f"Survey response rate is {len(survey)} of {len(students)} students.")
    target=number(fields.get("target"),2.4); direct=(number(fields.get("cie_weight"),60),number(fields.get("see_weight"),40)); overall=(number(fields.get("direct_weight"),80),number(fields.get("indirect_weight"),20))
    if not math.isclose(sum(direct),100,abs_tol=0.01): raise ValueError("CIE and SEE weights must total 100%.")
    if not math.isclose(sum(overall),100,abs_tol=0.01): raise ValueError("Direct and indirect weights must total 100%.")
    if not 0 <= target <= 3: raise ValueError("The attainment target must be between 0 and 3.")
    mapping=json.loads(fields.get("mapping",json.dumps(DEFAULT_MAPPING)))
    if len(mapping)!=4 or any(len(row)!=6 for row in mapping) or any(number(v,-1)<0 or number(v,-1)>3 for row in mapping for v in row): raise ValueError("CO–PO mapping must be a 4 × 6 matrix with values from 0 to 3.")
    known=set(registrations)
    unmatched=[x["reg"] for x in survey if x["reg"] not in known]
    if unmatched: warnings.append("Survey registration numbers not found in marks: "+", ".join(unmatched))
    for student in students:
        for label,marks,maxima in [("IA",student.ia,ia_max),("Mid-semester",student.mid,mid_max),("SEE",student.see,see_max)]:
            if any(mark<0 or mark>maxima[i] for i,mark in enumerate(marks)): raise ValueError(f"{label} marks for {student.reg} fall outside the permitted CO maximum marks.")
    result=compute(students,ia_max,mid_max,see_max,survey,direct,overall,mapping,target)
    output_dir.mkdir(parents=True,exist_ok=True); label=course.code.replace(' ','_') or 'Course'; name=f"{uuid.uuid4().hex[:10]}_CO_PO_Attainment_{label}.xlsx"; output=output_dir/name
    fill_template(files["template"],output,course,students,(ia_max,mid_max,see_max),survey,result,mapping,target,direct,overall)
    summary={"file":name,"students":len(students),"survey_responses":len(survey),"warnings":warnings,"cie":[round(x,3) for x in result["cie"]],"see":[round(x,3) for x in result["see"]],"direct":[round(x,3) for x in result["direct"]],"indirect":[round(x,3) for x in result["indirect"]] if survey else ["NA"]*4,"overall":[round(x,3) for x in result["overall"]],"po":[round(x,3) for x in result["overall_po"]],"target":target,"co_status":["YES" if x>=target else "NO" for x in result["overall"]]}
    return output,summary
